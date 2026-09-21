# -*- coding: utf-8 -*-
"""
Genera la carga inicial de la base de datos de la app de propinas.

Fuentes (todas ya existen en VARIABLES MENSUALES):
  1. propinas_la_rusia.html  -> la version original de la app, que lleva los
     datos incrustados (oct-2025 a mar-2026) y la lista de usuarios.
  2. <MES> <ANIO>/INFORME PERSONAL <MES> <ANIO>.xlsx -> un mes cada uno,
     incluido el consumo del periodo.

Salida:  carga_inicial.json
  Se importa UNA sola vez desde la consola de Firebase:
     Realtime Database -> menu (3 puntos) -> Importar JSON -> en el NODO RAIZ.

Las claves NO viajan en texto plano: se guarda SHA-256("RUT:clave").
La clave inicial de cada trabajador es su RUT sin puntos ni guion.
Las claves de la administracion viven solo en Firebase Authentication.

Uso:
    python generar_carga_inicial.py [carpeta_variables_mensuales]
"""
import hashlib
import json
import os
import re
import sys
import unicodedata

MESES = {'ENERO': 1, 'FEBRERO': 2, 'MARZO': 3, 'ABRIL': 4, 'MAYO': 5,
         'JUNIO': 6, 'JULIO': 7, 'AGOSTO': 8, 'SEPTIEMBRE': 9,
         'OCTUBRE': 10, 'NOVIEMBRE': 11, 'DICIEMBRE': 12}

# Correo de Firebase Authentication de cada administrador.
# Debe coincidir con la lista de database.rules.json.
ADMIN_EMAILS = {
    '17625819-5': 'oscar@larusia.app',
    '15610735-2': 'rossana@larusia.app',
    '15036876-6': 'paola@larusia.app',
}

# El convenio de propinas empieza en octubre de 2025: antes de esa fecha no
# hay fondo del 6% que acumular, asi que esos meses no se cargan.
MES_MINIMO = '2025-10'

CARPETA_DEFECTO = (r'C:\Users\oscar\Mi unidad (oscar.munizaga@gmail.com)\PERSONAL\GEST'
                   r'\CONTABILIDAD GEST\1.- CLIENTES ACTIVOS\LA RUSIA\RECURSOS HUMANOS'
                   r'\VARIABLES MENSUALES')


def sin_tildes(s):
    return ''.join(c for c in unicodedata.normalize('NFD', str(s))
                   if unicodedata.category(c) != 'Mn')


def norm_rut(v):
    s = re.sub(r'[^0-9kK]', '', str(v or ''))
    return s[:-1] + '-' + s[-1].upper() if len(s) >= 7 else None


def rut_display(r):
    cuerpo, dv = r.split('-')
    out = ''
    for i, c in enumerate(reversed(cuerpo)):
        if i and i % 3 == 0:
            out = '.' + out
        out = c + out
    return out + '-' + dv


def clave_inicial(rut):
    return rut.replace('.', '').replace('-', '')


def hash_pass(rut, clave):
    return hashlib.sha256((rut.upper() + ':' + clave).encode('utf-8')).hexdigest()


def num(v):
    if v is None or v == '':
        return 0
    if isinstance(v, (int, float)):
        return int(round(v))
    s = re.sub(r'[^\d,.\-]', '', str(v)).replace('.', '').replace(',', '.')
    try:
        return int(round(float(s)))
    except ValueError:
        return 0


# --------------------------------------------------------------------------
def leer_app_original(ruta):
    """Extrae el APPDATA incrustado en la version original de la app."""
    if not os.path.exists(ruta):
        print('   (no esta propinas_la_rusia.html: se parte de cero)')
        return {}, {}
    with open(ruta, encoding='utf-8', errors='replace') as fh:
        s = fh.read()
    i = s.find('const APPDATA = ')
    if i < 0:
        print('   (no se encontro APPDATA en la app original)')
        return {}, {}
    fin = s.find('\n', i)
    bruto = s[i + len('const APPDATA = '):fin].rstrip().rstrip(';')
    d = json.loads(bruto)
    workers, users = d.get('workers', {}), d.get('users', {})
    meses = sorted({h['m'] for w in workers.values() for h in w.get('h', [])})
    print('   app original: %d trabajadores, %d usuarios, %s..%s'
          % (len(workers), len(users), meses[0] if meses else '-',
             meses[-1] if meses else '-'))
    return workers, users


NOMBRE_MES = {v: k for k, v in MESES.items()}


def elegir_hoja(hojas, mes):
    """Hoja del mes que corresponde.

    Algunos INFORME PERSONAL (abril 2026, por ejemplo) traen una hoja por mes
    desde 2024; la primera hoja NO es la del mes de la carpeta.
    """
    anio, mm = mes.split('-')
    nombre = NOMBRE_MES[int(mm)]
    norm = [(h, sin_tildes(h).upper()) for h in hojas]
    for h, u in norm:
        if nombre in u and (anio in u or anio[-2:] in u):
            return h
    for h, u in norm:
        if nombre in u:
            return h
    return hojas[0]


def elegir_informe(carpeta, anio):
    """El INFORME PERSONAL del anio de la carpeta.

    Ojo: varias carpetas de 2026 conservan un 'INFORME PERSONAL 2025.xlsx'
    antiguo, y la de abril 2026 se llama solo 'INFORME PERSONAL.xlsx'.
    Por eso se prefiere el que nombra el anio y, si no hay, se acepta el que
    no nombra ningun anio; nunca uno que nombre OTRO anio.
    """
    candidatos = [f for f in os.listdir(carpeta)
                  if f.upper().startswith('INFORME PERSONAL')
                  and f.lower().endswith('.xlsx')
                  and not f.startswith('~$')]
    con_anio = [f for f in candidatos if anio in f or anio[-2:] in f]
    if con_anio:
        return sorted(con_anio, key=len)[0]
    sin_anio = [f for f in candidatos if not re.search(r'\d{4}|\d{2}', f)]
    return sin_anio[0] if sin_anio else None


def leer_informe_personal(ruta, mes):
    """Devuelve {rut: {n, p, r, net, c, consumo}} de un INFORME PERSONAL."""
    import openpyxl
    wb = openpyxl.load_workbook(ruta, data_only=True)
    ws = wb[elegir_hoja(wb.sheetnames, mes)]
    filas = list(ws.iter_rows(values_only=True))

    idx, enc = None, None
    for i, f in enumerate(filas[:8]):
        txt = [sin_tildes(c).upper().strip() if isinstance(c, str) else '' for c in f]
        if any('RUT' in t for t in txt) and any('PROPINA' in t for t in txt):
            idx, enc = i, txt
            break
    if idx is None:
        return {}

    def col(*claves, **kw):
        excluir = kw.get('excluir', ())
        for i, t in enumerate(enc):
            if t and not any(e in t for e in excluir) and all(k in t for k in claves):
                return i
        return None

    c_rut = col('RUT')
    c_nom = col('TRABAJADOR') if col('TRABAJADOR') is not None else col('NOMBRE')
    c_prop = col('PROPINAS', excluir=('NETA', 'BANCO', 'COMPENS'))
    c_ret = col('6%')
    c_net = col('NETAS') if col('NETAS') is not None else col('A PAGAR')
    c_comp = col('COMPENS')
    c_cons = col('CONSUMO')
    if c_rut is None or c_prop is None:
        return {}

    def cel(f, c):
        return num(f[c]) if c is not None and c < len(f) else 0

    out = {}
    for f in filas[idx + 1:]:
        rut = norm_rut(f[c_rut] if c_rut < len(f) else None)
        if not rut:
            continue
        p = cel(f, c_prop)
        r = cel(f, c_ret) or round(p * 0.06)
        net = cel(f, c_net) or max(0, p - r)
        cons = cel(f, c_cons)
        if not (p or net or cons):
            continue
        out[rut] = {
            'n': str(f[c_nom] or '').strip() if c_nom is not None and c_nom < len(f) else '',
            'p': p, 'r': r, 'net': net, 'c': cel(f, c_comp), 'consumo': cons,
        }
    return out


# --------------------------------------------------------------------------
def main():
    base = sys.argv[1] if len(sys.argv) > 1 else CARPETA_DEFECTO
    if not os.path.isdir(base):
        print('No existe la carpeta: %s' % base)
        return 1

    print('Leyendo %s' % base)
    workers, users_orig = leer_app_original(os.path.join(base, 'propinas_la_rusia.html'))
    workers = {norm_rut(k) or k: v for k, v in workers.items()}
    meses_base = {h['m'] for w in workers.values() for h in w.get('h', [])}

    consumo = {}
    nuevos_meses = []
    for carpeta in sorted(os.listdir(base)):
        ruta_c = os.path.join(base, carpeta)
        if not os.path.isdir(ruta_c):
            continue
        m = re.match(r'^([A-Z]+)\s+(\d{2,4})$', sin_tildes(carpeta).upper().strip())
        if not m or m.group(1) not in MESES:
            continue
        anio = m.group(2)
        anio = '20' + anio if len(anio) == 2 else anio
        mes = '%s-%02d' % (anio, MESES[m.group(1)])
        if mes < MES_MINIMO:
            continue          # anterior al convenio de propinas
        if mes in meses_base:
            continue          # ese mes ya viene de la app original (dato bueno)
        archivo = elegir_informe(ruta_c, anio)
        if not archivo:
            continue
        datos = leer_informe_personal(os.path.join(ruta_c, archivo), mes)
        if not datos:
            continue
        nuevos_meses.append((mes, len(datos)))
        for rut, v in datos.items():
            w = workers.setdefault(rut, {'n': v['n'] or rut,
                                         'rd': rut_display(rut), 'h': []})
            if v['n']:
                w['n'] = w.get('n') or v['n']
            w['h'] = [h for h in w.get('h', []) if h['m'] != mes]
            w['h'].append({'m': mes, 'p': v['p'], 'r': v['r'], 'net': v['net'],
                           'c': v['c'], 'nf': v['net'] + v['c'], 's': 0})
            if v['consumo']:
                c = consumo.setdefault(rut, {'n': w['n'], 'h': []})
                c['h'] = [x for x in c['h'] if x['m'] != mes]
                c['h'].append({'m': mes,
                               'items': [{'det': 'Consumo del mes', 'monto': v['consumo']}],
                               'total': v['consumo']})

    # Recalcular saldo acumulado del fondo 6% y ordenar
    for w in workers.values():
        w['h'] = sorted(w.get('h', []), key=lambda h: h['m'])
        saldo = 0
        for h in w['h']:
            for k in ('p', 'r', 'net', 'c', 'nf'):
                h[k] = int(round(h.get(k) or 0))
            saldo += h['r'] - h['c']
            h['s'] = saldo
            h['nf'] = h['net'] + h['c']
    for c in consumo.values():
        c['h'] = sorted(c['h'], key=lambda x: x['m'])

    # Usuarios: trabajadores con hash; administracion contra Firebase Auth
    users = {}
    for rut, w in workers.items():
        orig = users_orig.get(rut) or users_orig.get(rut.replace('-', '')) or {}
        nombre = orig.get('nombre') or w['n']
        if rut in ADMIN_EMAILS or orig.get('tipo') == 'admin':
            users[rut] = {'nombre': nombre, 'tipo': 'admin', 'changed': True,
                          'email': ADMIN_EMAILS.get(rut, 'oscar@larusia.app')}
        else:
            users[rut] = {'nombre': nombre, 'tipo': 'trabajador', 'changed': False,
                          'passHash': hash_pass(rut, clave_inicial(rut))}
    # Administradores que no son trabajadores
    for rut, correo in ADMIN_EMAILS.items():
        if rut not in users:
            orig = users_orig.get(rut, {})
            users[rut] = {'nombre': orig.get('nombre', rut), 'tipo': 'admin',
                          'changed': True, 'email': correo}

    salida = {'propinas': {'workers': workers, 'users': users,
                           'consumo': consumo, 'config': {'admins': {}}}}
    destino = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'carga_inicial.json')
    with open(destino, 'w', encoding='utf-8') as fh:
        json.dump(salida, fh, ensure_ascii=False, indent=1)

    meses = sorted({h['m'] for w in workers.values() for h in w['h']})
    total = sum(h['p'] for w in workers.values() for h in w['h'])
    print('\nMeses agregados desde los INFORME PERSONAL:')
    for mes, n in nuevos_meses:
        print('   %s  (%d trabajadores)' % (mes, n))
    print('\nListo: %s' % destino)
    print('   %d trabajadores | %d usuarios (%d admin) | %d con consumo'
          % (len(workers), len(users),
             sum(1 for u in users.values() if u['tipo'] == 'admin'), len(consumo)))
    print('   meses: %s .. %s' % (meses[0], meses[-1]))
    print('   total propinas historicas: $%s' % format(total, ',d').replace(',', '.'))
    print('\nImportar en: consola de Firebase -> Realtime Database -> Importar JSON (nodo raiz)')
    return 0


if __name__ == '__main__':
    sys.exit(main())
